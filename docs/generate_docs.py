from pathlib import Path
import csv
import zipfile

from fpdf import FPDF
from docx import Document

ROOT = Path(r"C:\Users\Kevan\OneDrive - FTH Trading\11-Downloads\unykorn-black-funder-site")
DOWNLOADS = ROOT / "downloads"
DOWNLOADS.mkdir(parents=True, exist_ok=True)

SCENARIOS = [
    ("Lean", "5 vehicles", "$723,150", "2.6x"),
    ("Base", "6 vehicles", "$1,414,200", "7.2x"),
    ("Growth", "8 vehicles", "$2,365,000", "10.5x"),
]

USE_OF_FUNDS = [
    ("Vehicle acquisition / down payments", "$115,000"),
    ("Equipment / vehicle financing facility", "$350,000"),
    ("Insurance reserve", "$45,000"),
    ("Licensing, compliance, legal, accounting", "$25,000"),
    ("Website, app, dispatch buildout", "$35,000"),
    ("Branding, QR, sponsor materials", "$20,000"),
    ("Working capital reserve", "$60,000"),
]

EVENTS = [
    ("May 24, 2026", "Birthday Bash ATL", "State Farm Arena", "B+", "Late-night VIP routes, after-event offers, sponsor nightlife traffic"),
    ("June 8-9, 2026", "Megan Moroney", "State Farm Arena", "B", "Concert transfers, hotel pickup, premium downtown routes"),
    ("June 9-14, 2026", "Atlanta Market + Atlanta Apparel", "AmericasMart", "A", "Buyers, vendors, showroom executives, airport and hotel routes"),
    ("June 15-July 15, 2026", "Atlanta World Cup match window", "Atlanta Stadium / Mercedes-Benz Stadium", "A+", "VIP arrivals, hotel-stadium-dining loops, sponsors"),
    ("Aug. 21, 2026", "Chris Stapleton", "Mercedes-Benz Stadium", "A", "Stadium concert, suite guests, sponsor hospitality, reserved exits"),
    ("Aug. 26-30, 2026", "TOUR Championship", "East Lake Golf Club", "A", "Executive golf, sponsor hospitality, corporate blocks, airport transfers"),
    ("Aug. 27, 2026", "AC/DC POWER UP Tour", "Mercedes-Benz Stadium", "A", "Major stadium concert, premium arrivals and late-night exits"),
    ("Sept. 3-7, 2026", "Dragon Con", "Downtown Atlanta", "A", "Multi-day hotel corridor, fan groups, VIP events, merchant offers"),
    ("Sept. 18-20, 2026", "Shaky Knees", "Piedmont Park", "A", "Festival rides, hotel packages, sponsor QR offers, group exits"),
    ("Oct. 10-11, 2026", "Atlanta Pride", "Piedmont Park", "B+", "Festival mobility, hotel routes, merchant and safety-oriented routing"),
    ("Nov. 10-11, 2026", "The R&B Tour", "Mercedes-Benz Stadium", "A", "High-demand concert nights, suite guests, sponsors, premium exits"),
    ("Dec. 5, 2026", "SEC Championship", "Mercedes-Benz Stadium", "A+", "Corporate hospitality, alumni groups, VIP fan travel, sponsor routes"),
]


def add_title(pdf: FPDF, title: str, subtitle: str) -> None:
    pdf.set_fill_color(7, 17, 31)
    pdf.set_text_color(214, 168, 79)
    pdf.set_font("Helvetica", "B", 19)
    pdf.cell(0, 12, title, ln=True)
    pdf.set_text_color(40, 40, 40)
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, subtitle)
    pdf.ln(3)


def build_lender_pdf() -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    add_title(
        pdf,
        "UNYKORN BLACK - Lender Proposal",
        "Premium Event Mobility for Atlanta 2026. Requested facility: $650,000 for a 5-8 vehicle launch and sponsor-backed growth.",
    )

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Executive Summary", ln=True)
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(
        0,
        6,
        "UNYKORN BLACK is a premium event transportation and media platform. It combines luxury fleet operations with sponsor route inventory, merchant referrals, and TROPTIONS Pay settlement reporting. The model is designed to lower risk through prepaid ride credits (BlackPass), account-backed dispatch, and partner-fleet overflow.",
    )
    pdf.ln(2)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Financial Scenarios", ln=True)
    pdf.set_font("Helvetica", "", 11)
    for name, fleet, rev, dscr in SCENARIOS:
        pdf.cell(0, 7, f"{name}: {fleet} | Revenue {rev} | DSCR {dscr}", ln=True)

    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Use of Funds", ln=True)
    pdf.set_font("Helvetica", "", 11)
    for item, amount in USE_OF_FUNDS:
        pdf.cell(140, 7, item)
        pdf.cell(0, 7, amount, ln=True)

    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Collateral and Risk Controls", ln=True)
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(
        0,
        6,
        "1) Vehicle collateral on core fleet units. 2) Prepaid BlackPass revenue and sponsor contracts. 3) Partner fleet overflow to avoid overbuying. 4) Route-level settlement visibility through TROPTIONS Pay reporting. 5) Launch phase controls tied to event demand windows.",
    )

    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Atlanta Event Demand Calendar", ln=True)
    pdf.set_font("Helvetica", "", 10)
    for row in EVENTS:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(190, 5.5, f"{row[0]} | {row[1]} | {row[2]} | Tier {row[3]} | {row[4]}")

    pdf.output(str(DOWNLOADS / "unykorn-black-lender-proposal.pdf"))


def build_rate_card_pdf() -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    add_title(
        pdf,
        "UNYKORN BLACK - Sponsor Rate Card",
        "Premium inventory across vehicles, routes, app flows, and merchant activations.",
    )
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(110, 8, "Package")
    pdf.cell(35, 8, "Price")
    pdf.cell(0, 8, "Description", ln=True)
    pdf.set_font("Helvetica", "", 11)

    rows = [
        ("Founding Sponsor", "$150,000", "Top placement across site, app, fleet, receipts, proof reports"),
        ("Route Sponsor", "$25,000", "Exclusive airport-hotel-stadium or post-event corridor ownership"),
        ("Vehicle Sponsor", "$15,000", "Branded SUV/Sprinter placement with QR and confirmation visibility"),
        ("Merchant Listing", "$2,500", "Verified premium listing with referral and offer placement"),
    ]

    for package, price, desc in rows:
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(190, 6.5, f"{package} | {price} | {desc}")

    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Contact", ln=True)
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, "UNYKORN BLACK Partnerships\nWebsite: https://fifa.unykorn.org/\nPurpose: lender + sponsor launch package")

    pdf.output(str(DOWNLOADS / "unykorn-black-sponsor-rate-card.pdf"))


def build_docx() -> None:
    doc = Document()
    doc.add_heading("UNYKORN BLACK - Editable Funding Proposal", level=1)
    doc.add_paragraph(
        "Premium event mobility launch package for Atlanta 2026. This editable file is prepared for lender and partner customization."
    )

    doc.add_heading("Funding Request", level=2)
    doc.add_paragraph("$650,000 controlled launch package for a 5-8 vehicle premium event fleet.")

    doc.add_heading("Scenario Snapshot", level=2)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Scenario"
    hdr[1].text = "Fleet"
    hdr[2].text = "Gross Revenue"
    hdr[3].text = "DSCR"
    for row in SCENARIOS:
        c = table.add_row().cells
        c[0].text = row[0]
        c[1].text = row[1]
        c[2].text = row[2]
        c[3].text = row[3]

    doc.add_heading("Use of Funds", level=2)
    for item, amount in USE_OF_FUNDS:
        doc.add_paragraph(f"{item}: {amount}", style="List Bullet")

    doc.add_heading("Event Demand Window", level=2)
    for date, event, venue, tier, angle in EVENTS:
        doc.add_paragraph(f"{date} | {event} | {venue} | Tier {tier} | {angle}")

    doc.add_paragraph("Powered by TROPTIONS reporting and WhichWay platform integration.")

    doc.save(str(DOWNLOADS / "unykorn-black-lender-proposal.docx"))


def build_csv() -> None:
    csv_path = DOWNLOADS / "unykorn-black-event-demand-calendar.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["date", "event", "venue", "tier", "revenue_angle"])
        writer.writerows(EVENTS)


def build_zip() -> None:
    zip_path = DOWNLOADS / "unykorn-black-downloads.zip"
    members = [
        "unykorn-black-lender-proposal.pdf",
        "unykorn-black-sponsor-rate-card.pdf",
        "unykorn-black-lender-proposal.docx",
        "unykorn-black-event-demand-calendar.csv",
    ]
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for m in members:
            zf.write(DOWNLOADS / m, arcname=m)


def main() -> None:
    build_lender_pdf()
    build_rate_card_pdf()
    build_docx()
    build_csv()
    build_zip()


if __name__ == "__main__":
    main()
